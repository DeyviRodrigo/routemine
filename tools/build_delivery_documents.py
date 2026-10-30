"""Build the RouteTec manuals and a complete, verifiable source-code volume.

Run with the document runtime supplied by Codex or with python-docx and Pillow.
Screenshots are real application captures, never drawn UI mockups.
"""
from __future__ import annotations
import argparse
import hashlib
import json
import textwrap
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'entrega' / 'RouteTec'
QA = ROOT / 'out' / 'delivery-docs'
BLUE = '18556A'

def field(p, instruction):
    f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), instruction); p._p.append(f)

def setup(title, subtitle, landscape=False):
    d = Document(); s=d.sections[0]
    s.page_width=Inches(11 if landscape else 8.5)
    s.page_height=Inches(8.5 if landscape else 11)
    s.top_margin=s.bottom_margin=Inches(.65)
    s.left_margin=s.right_margin=Inches(.7)
    normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
    normal.font.color.rgb=RGBColor(0,0,0)
    normal.paragraph_format.space_after=Pt(7)
    normal.paragraph_format.line_spacing=1.08
    for name,size in [('Title',28),('Heading 1',19),('Heading 2',14),('Heading 3',12)]:
        style=d.styles[name];style.font.name='Calibri';style.font.size=Pt(size)
        style.font.color.rgb=RGBColor.from_string('000000' if name=='Title' else BLUE)
        style.paragraph_format.space_before=Pt(13);style.paragraph_format.space_after=Pt(8)
    for style in d.styles:
        if style._element.pPr is not None:
            for border in list(style._element.pPr.findall(qn('w:pBdr'))):
                style._element.pPr.remove(border)
    d.styles['Subtitle'].font.color.rgb=RGBColor.from_string('555555')
    h=s.header.paragraphs[0];h.text='RouteTec    |    Versión 1.0.0';h.style='Caption'
    foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    foot.add_run('RouteTec   ·   ');field(foot,'PAGE')
    d.add_paragraph(title,'Title');d.add_paragraph(subtitle,'Subtitle')
    d.add_paragraph('Edición del 30 de octubre de 2026')
    return d

def body(d, text):
    return d.add_paragraph(text)

def table(d, headers, rows):
    t=d.add_table(rows=1, cols=len(headers));t.style='Light Shading Accent 1'
    if len(headers)==2:
        t.autofit=False;t.columns[0].width=Inches(2.15);t.columns[1].width=Inches(4.95)
    for c,v in zip(t.rows[0].cells,headers):c.text=v
    pr=t.rows[0]._tr.get_or_add_trPr();pr.append(OxmlElement('w:tblHeader'))
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for i,row in enumerate(t.rows):
        if len(headers)==2:
            row.cells[0].width=Inches(2.15);row.cells[1].width=Inches(4.95)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(4)
                for r in p.runs:r.font.color.rgb=RGBColor.from_string(BLUE if i==0 else '111111')
    return t

def figure(d, path, caption, max_height=5.6):
    path=Path(path)
    if not path.exists():raise FileNotFoundError(path)
    w,h=Image.open(path).size
    width=min(7.05,max_height*w/h)
    p=d.add_paragraph();p.paragraph_format.keep_with_next=True
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run();r.add_picture(str(path),width=Inches(width))
    inline=r._r.xpath('.//wp:docPr')
    if inline:inline[0].set('descr',caption)
    p=d.add_paragraph(caption,'Caption');p.paragraph_format.keep_with_next=False

def chapter(d,title,text=None):
    heading=d.add_heading(title,1);heading.paragraph_format.page_break_before=True
    if text:body(d,text)

def source_files():
    extensions={'.cpp','.hpp','.h','.cu','.c','.py','.ps1','.bat','.cmake','.nsi','.iss','.json','.yaml','.yml','.txt','.request','.md','.rc','.manifest','.in'}
    fs=[]
    for folder in ['app','include','src','tests','packaging','tools']:
        base=ROOT/folder
        if base.exists():
            fs += [p for p in base.rglob('*') if p.is_file() and p.suffix.lower() in extensions and '__pycache__' not in p.parts]
    fs += [p for p in ROOT.iterdir() if p.is_file() and (p.name=='CMakeLists.txt' or p.suffix in {'.bat','.yml','.yaml'})]
    return sorted(set(fs),key=lambda p:p.relative_to(ROOT).as_posix().casefold())

FOLDER_INFO={
    'app':'Interfaz nativa, visualización tridimensional, persistencia de proyectos y entrada de la línea de comandos.',
    'include':'Contratos públicos del núcleo de cálculo, modelos geométricos y configuración.',
    'include/routemine':'Cabeceras del espacio de nombres técnico conservado por compatibilidad con el núcleo.',
    'src':'Algoritmos de búsqueda, mallas, geometría, validación, movimientos de tierra y exportación.',
    'tests':'Pruebas automatizadas, regresiones geométricas y programas de evaluación con terrenos reales.',
    'packaging':'Preparación de dependencias, instalación y distribución para Windows.',
    'tools':'Herramientas de documentación y comprobación de la entrega.',
    'raíz':'Definición de compilación, restauración del entorno y lanzadores locales.'}

def describe(p):
    n=p.name
    exact={
      'viewer.cpp':'Ventana principal, diálogos de configuración, tareas asíncronas, controles y pruebas gráficas.',
      'terrain_view.cpp':'Escena VTK, terreno y carretera, cámaras, selección de puntos, leyenda y rutas exploradas.',
      'project_io.cpp':'Lectura y escritura atómica de proyectos JSON y restauración de rutas.',
      'core.hpp':'Tipos, parámetros y funciones públicas de geometría, búsqueda, cálculo y validación.',
      'config.hpp':'Interfaz pública de resolución y persistencia de configuración.',
      'config.cpp':'Valores iniciales, lectura de parámetros, validación y serialización de configuración.',
      'ramp_search.cpp':'Planificación de curvas continuas y reconstrucción de rasante validada.',
      'road_search.cpp':'Selección de estrategias de búsqueda y planificación de alternativas.',
      'adaptive_road.cpp':'Secuencia controlada de intentos de planificación y registro de parámetros efectivos.',
      'route_refinement.cpp':'Refinamiento y optimización de intervalos de una ruta.',
      'alignment.cpp':'Geometría y validación de alineaciones tangentes y curvas.',
      'geometry.cpp':'Consultas geométricas de la malla y perfiles sobre sus caras originales.',
      'validation.cpp':'Comprobación independiente de pendiente, cobertura y límites geométricos.',
      'cost.cpp':'Cantidades y objetivo de comparación de rutas.',
      'dxf.cpp':'Importación de caras de terreno desde DXF.',
      'obj.cpp':'Importación y validación de mallas OBJ.',
      'terrain_io.cpp':'Selección del lector según el formato del archivo topográfico.',
      'earthwork_segments.cpp':'Clasificación de segmentos de superficie, corte y relleno.',
      'route_export.cpp':'Exportación de ejes y segmentos a CSV y DXF.',
      'search.cpp':'Búsqueda discreta del planificador original.',
      'backend_cpu.cpp':'Backend de cálculo de pendientes en CPU.',
      'backend_cuda.cu':'Backend CUDA opcional que no se requiere para la edición distribuida.',
      'search_trace.hpp':'Intercambio acotado entre búsqueda y visualización de ramas exploradas.',
      'main.cpp':'Entrada de los comandos de inspección, demostración y planificación.',
      'plan_cli.cpp':'Ejecución de planificación por línea de comandos y escritura del resultado.',
      'CMakeLists.txt':'Definición de biblioteca, ejecutables, opciones, dependencias y pruebas.',
      'build_delivery_documents.py':'Generación de manuales ilustrados y del volumen completo del código fuente.'}
    if n in exact:return exact[n]
    if p.parent.name=='tests':return 'Casos y utilidades de prueba para '+p.stem.replace('_',' ')+'.'
    if p.suffix=='.hpp':return 'Declaraciones y contratos utilizados por '+p.stem.replace('_',' ')+'.'
    if p.suffix=='.bat':return 'Lanzador Windows del aplicativo o de un escenario de prueba.'
    if 'packaging' in p.parts:return 'Herramienta de distribución e instalación de RouteTec para Windows.'
    return 'Configuración o herramienta auxiliar de '+p.stem.replace('_',' ')+'.'

def build_source():
    d=setup('Código fuente de RouteTec','Organización por carpetas y archivos con transcripción íntegra del código',True)
    files=source_files()
    body(d,'Este volumen contiene el código propio del aplicativo, las pruebas y las herramientas de compilación, documentación y empaquetado. Los nombres técnicos del espacio de nombres routemine se conservan para compatibilidad; el nombre comercial y el ejecutable de distribución son RouteTec.')
    body(d,'No se transcriben bibliotecas externas Qt, VTK y runtimes, carpetas generadas out y build, historial .git, ejecutables, capturas ni documentos binarios. Las licencias de terceros se distribuyen con el instalador. Los terrenos de ejemplo son datos y no código fuente.')
    manifest=[]
    table(d,['Carpeta','Descripción'],list(FOLDER_INFO.items()))
    table(d,['Archivos','Integridad'],[(str(len(files)),'Cada archivo incluye SHA-256 y número de líneas para cotejar con el repositorio entregado.')])
    for p in files:
        raw=p.read_bytes(); text=raw.decode('utf-8-sig');lines=text.splitlines()
        rel=p.relative_to(ROOT).as_posix();sha=hashlib.sha256(raw).hexdigest()
        manifest.append({'path':rel,'sha256':sha,'lines':len(lines),'bytes':len(raw)})
        chapter(d,rel)
        body(d,describe(p));body(d,f'{len(lines)} líneas · {len(raw)} bytes\nSHA-256 {sha}')
        for line_index,line in enumerate(lines):
            par=d.add_paragraph();pf=par.paragraph_format
            pf.space_before=pf.space_after=Pt(0);pf.line_spacing=1
            # Keep the file's final block together so a closing brace cannot
            # become an otherwise empty continuation page.
            pf.keep_with_next=(line_index>=max(0,len(lines)-8) and line_index<len(lines)-1)
            pf.widow_control=False
            # Native Word soft wrapping retains every source character, including tabs.
            run=par.add_run(line);run.font.name='Consolas';run.font.size=Pt(9)
    target=OUT/'Código fuente.docx';d.save(target)
    QA.mkdir(parents=True,exist_ok=True)
    (QA/'source-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    return target

USER_PAGES = [
('01_inicio','Conocer la pantalla principal',
 'RouteTec prepara un eje de carretera entre un origen A y un destino B sobre una topografía tridimensional. Este manual explica el uso del aplicativo ya instalado. La guía de instalación se entrega por separado.',
 [('Terreno','Cargar una malla DXF u OBJ.'),('A y B','Registrar coordenadas o seleccionar cada extremo en la vista 3D.'),('Calcular ruta','Buscar un recorrido con el diseño configurado.'),('Resultados','Consultar las medidas de la ruta calculada. Los detalles se abren desde Ayuda.')]),
('26_cargar_topografia','Cargar la topografía',
 'Pulse Cargar topografía, entre en la carpeta de su levantamiento, seleccione un archivo DXF u OBJ y confirme Abrir. Espere a que termine la carga antes de marcar los extremos.',
 [('DXF','Debe contener caras 3DFACE del terreno; un plano con solo líneas o curvas de nivel no es una malla.'),('OBJ','Debe representar una superficie triangulada con cotas.'),('Unidades','Utilice metros en X, Y y Z. Las coordenadas se interpretan en el sistema del archivo; no se reproyectan.'),('Malla','Los huecos y las cotas ambiguas no se rellenan automáticamente.')]),
('05_panel_controles','Registrar el origen y el destino',
 'Para introducir coordenadas, complete X, Y y Z y pulse Registrar origen o Registrar destino. Escribir valores sin pulsar Registrar no confirma el punto. Use punto decimal y evite separadores de miles.',
 [('Seleccionar origen en 3D','Pulse el botón y después haga clic sobre la superficie.'),('Seleccionar destino en 3D','Repita la operación para B; RouteTec obtiene la cota del terreno.'),('Coordenadas','Compruebe los tres valores y que A y B sean diferentes.'),('Limpiar ruta y puntos','Retira la carretera y los extremos para comenzar otro recorrido; conserva la topografía.')]),
('04_ruta_iso','Elegir el tipo de propuesta',
 'El selector de propuesta decide qué se calculará. Seleccione el modo antes de pulsar Calcular ruta. Si ya calculó ambos modos, puede alternar sus resultados.',
 [('Propuesta respetando el terreno','El eje sigue las cotas originales de la malla. No proyecta corte ni relleno.'),('Propuesta con corte y relleno','La rasante puede separarse de la malla dentro de los límites configurados.'),('Calcular ruta','Inicia una búsqueda nueva con los puntos y parámetros vigentes.'),('Comparación','Un resultado del segundo modo puede ser más corto o no lograr mejorar el primero; no hay garantía de óptimo global.')]),
('09_menu_configuracion','Abrir la configuración',
 'El menú Configuración reúne el diseño general, los ajustes avanzados y los puntos intermedios. Así la pantalla principal conserva solo las acciones de trabajo.',
 [('Diseño','Pendiente, radio y recta de transición al invertir el sentido de giro.'),('Configuración avanzada','Corte, relleno, ancho, precisión, recursos y costos.'),('Puntos intermedios','Habilitación de las cuatro acciones laterales y parámetros del tramo local.'),('Recalcular','Si modifica el diseño general, vuelva a calcular la carretera para obtener un resultado con esos valores.')]),
('10_diseno','Configurar la geometría de la carretera',
 'Los valores iniciales del diseño son pendiente máxima 10 %, radio mínimo 30 m y recta mínima de 35 m al invertir el giro. Mantenga estos valores para el primer intento y cambie únicamente lo que requiera su proyecto.',
 [('Pendiente máxima','Porcentaje de subida o bajada máxima del eje. Una pendiente 0 % permite únicamente recorridos horizontales.'),('Radio mínimo','Límite inferior del radio de curva en planta; un radio menor facilita rodeos cerrados.'),('Recta al invertir el giro','Longitud mínima entre un giro a la izquierda y otro a la derecha, o a la inversa. No obliga a interrumpir curvas del mismo sentido.'),('Extender los límites','Amplía el rango editable. No mejora por sí solo la búsqueda ni certifica un diseño. Desactivarlo ajusta los valores al rango habitual.')]),
('11_avanzado_diseno','Definir corte relleno y ancho',
 'En Configuración avanzada, pestaña Diseño, establezca cuánto puede separarse verticalmente el eje del terreno. Estos límites se usan en la propuesta con corte y relleno.',
 [('Corte máximo','Profundidad máxima bajo la superficie original, en metros. Inicial 2 m; cero prohíbe el corte.'),('Relleno máximo','Altura máxima sobre la superficie original, en metros. Inicial 2 m; cero prohíbe el relleno.'),('Ancho de cálculo','Se usa en cantidades preliminares y separación entre vueltas. Inicial 8 m.'),('Restablecer valores','Restaura los ajustes avanzados y geométricos iniciales, conservando terreno, extremos, pendiente máxima y modo seleccionado. Debe recalcular.')]),
('12_avanzado_precision','Ajustar la precisión',
 'La búsqueda genera candidatos discretos y verifica el resultado sobre las caras originales del terreno. Use primero la precisión inicial; reducir todos los pasos a la vez incrementa el trabajo y puede agotar los recursos.',
 [('Paso XY objetivo','Escala espacial de exploración. Inicial 10 m; el buscador puede ensayar pasos efectivos distintos.'),('Muestreo','Separación objetivo para reconstruir y verificar el recorrido. Inicial 2 m; los cruces de caras añaden muestras.'),('Tolerancia de superficie','Margen numérico al comparar el eje con la malla. Inicial 0,01 m; no es un permiso de corte de obra.'),('Paso de direcciones','Resolución angular de exploración. Inicial 7,5 grados. La búsqueda puede complementar direcciones en laderas.')]),
('13_avanzado_recursos','Controlar los recursos de búsqueda',
 'Los recursos limitan memoria y trabajo. Llegar al límite no demuestra que el recorrido sea imposible. Mantenga activada la adaptación para permitir los refinamientos internos previstos.',
 [('Adaptar resolución y recursos','Permite variantes de paso y presupuestos, manteniendo el diseño geométrico.'),('Nodos y expansiones','Límites iniciales de estados almacenados y estados procesados.'),('Aristas','Límite de transiciones comprobadas durante la búsqueda.'),('Muestras y candidatos','Acotan la reconstrucción del recorrido y el número de candidatos completos evaluados por variante.')]),
('14_avanzado_costos','Elegir el objetivo de comparación',
 'La pestaña Costos permite comparar por longitud o por tarifas configuradas. El perfil inicial Longitud no representa un presupuesto monetario.',
 [('Longitud sin precios','Compara recorridos usando pesos de longitud. Las tarifas de corte y relleno no intervienen como precios reales.'),('Costos configurados','Habilita tarifas y distancias para una estimación preliminar coherente con los datos que introduzca.'),('Ejemplo sintético','Contiene valores ficticios útiles para pruebas. No los use como cotización.'),('Unidad de costo','Todas las tarifas deben emplear la misma moneda o unidad de costo, con las unidades mostradas en cada campo.')]),
('02_progreso_carga','Supervisar una operación',
 'La ventana de progreso aparece durante la carga y el cálculo. Muestra la etapa, el porcentaje, el tiempo transcurrido y una estimación del tiempo restante.',
 [('Porcentaje','Describe el avance de la operación; distintas etapas pueden tener velocidades diferentes.'),('Tiempo restante','Es aproximado y puede cambiar al pasar a otra etapa.'),('Minimizar','Oculta la ventana flotante y conserva el proceso en marcha.'),('Cancelar','Solicita detener el trabajo; espere a que se confirme la cancelación antes de iniciar otro.')]),
('03_progreso_minimizado','Seguir trabajando con el progreso minimizado',
 'Al minimizar, aparece una franja de progreso dentro de la ventana principal. Cerrar la ventana flotante o pulsar Escape también la minimiza; no equivale a cancelar.',
 [('Ver progreso','Vuelve a mostrar la ventana flotante.'),('Cancelar','Detiene la operación desde la franja inferior.'),('Finalización','La franja se retira cuando la operación concluye.'),('Durante el cálculo','Los parámetros y las operaciones que cambian el proyecto se bloquean para conservar un estado coherente.')]),
('07_leyenda_y_ruta','Leer los colores de las rutas',
 'La leyenda está dentro de la vista tridimensional. Los colores se refieren al eje de la carretera y a la exploración; la escala de cotas del terreno es independiente.',
 [('Azul sólido','Ruta calculada o tramo sobre el terreno.'),('Rojo','Tramo clasificado como corte.'),('Verde','Tramo clasificado como relleno.'),('Gris y azul transparentes','Ramas exploradas y candidatos completos. Se pueden ocultar desde Vistas. P1 y P2 identifican el intervalo local.')]),
('06_ruta_superior','Revisar la carretera en planta',
 'La vista superior permite seguir los rodeos, las curvas y la relación entre origen y destino. Revise también una vista oblicua para entender el desnivel.',
 [('Longitud','Desarrollo horizontal y distancia espacial, cuando se muestran en los detalles.'),('Pendiente','Máxima pendiente comprobada del eje calculado.'),('Vueltas','Las curvas consecutivas del mismo sentido permiten rodear el terreno.'),('Extremos','La ruta debe comenzar en A y terminar en B. Cambiar un extremo exige otro cálculo.')]),
('19_menu_vistas','Cambiar la vista y la exploración',
 'El menú Vistas modifica la cámara sin alterar los datos del terreno o el diseño de la carretera.',
 [('ISO y Superior','Perspectiva oblicua y planta.'),('Inferior y laterales','Permiten revisar el terreno desde abajo, frente, atrás, izquierda o derecha.'),('Acercar y alejar','Ajustan la escala de visualización.'),('Mostrar caminos explorados','Activa u oculta las ramas de búsqueda; no elimina el resultado ni cambia su cálculo.')]),
('22_vista_frontal','Manipular la escena tridimensional',
 'Use el ratón sobre la escena para inspeccionar el terreno y la carretera. Si va a seleccionar un punto, pulse primero su botón de selección para que el clic tenga ese significado.',
 [('Arrastre izquierdo','Gira la cámara con el comportamiento de navegación de la escena.'),('Rueda','Acerca o aleja la vista.'),('Desplazamiento','Use el botón central o la combinación de desplazamiento de la cámara VTK.'),('Recuperar orientación','Seleccione ISO o Superior en Vistas cuando necesite volver a una vista conocida.')]),
('15_configuracion_intermedios','Habilitar los puntos intermedios',
 'Abra Configuración y Puntos intermedios. Marque Habilitar puntos intermedios y pulse Aplicar. El lateral muestra las cuatro acciones en dos filas.',
 [('Habilitar','Muestra u oculta las acciones laterales sin cambiar la geometría de la ruta por sí solo.'),('P1 y P2','Seleccionan el primer y último punto del intervalo; P1 debe estar antes que P2 en el orden del recorrido.'),('Parámetros locales','Pendiente, radio y recta se definen en este mismo diálogo.'),('Aplicar y Cancelar','Aplicar conserva la configuración. Cancelar cierra sin aplicar las ediciones de esta ventana.')]),
('17_cuatro_botones_intermedios','Usar las cuatro acciones del tramo',
 'Primero calcule una ruta. Después marque dos puntos sobre su eje, revise los parámetros locales y optimice el intervalo.',
 [('Marcar P1','Activa la selección del inicio del tramo; haga clic sobre la carretera.'),('Marcar P2','Activa la selección del final del tramo; haga clic más adelante en la carretera.'),('Optimizar tramo','Busca una alternativa dentro del intervalo respetando sus extremos y parámetros.'),('Limpiar','Retira la selección P1 y P2; no borra la topografía ni la ruta calculada.')]),
('18_parametros_tramo','Configurar las restricciones locales',
 'Los parámetros del tramo se muestran en Configuración y Puntos intermedios. El resto de la carretera conserva las restricciones con las que se calculó.',
 [('Pendiente máxima local','Límite de pendiente exclusivamente entre P1 y P2.'),('Radio mínimo local','Radio permitido para la alternativa local.'),('Recta local','Recta exigida al invertir el sentido del giro en ese intervalo.'),('Orden de trabajo','Marcar P1 y P2, configurar, aplicar y pulsar Optimizar tramo. Cambiar campos no optimiza automáticamente.')]),
('16_puntos_p1_p2','Revisar una optimización local',
 'Los marcadores muestran qué parte de la carretera está seleccionada. La optimización mantiene el recorrido exterior y comprueba las uniones del tramo con la carretera.',
 [('Resultado mejorado','Si se obtiene una alternativa válida que mejora el objetivo, se incorpora al recorrido.'),('Sin mejora','La ruta anterior se conserva cuando no se consigue una mejora válida.'),('Selección incorrecta','Si P1 queda después de P2, vuelva a marcar o corrija el orden en Configuración.'),('Guardar','Guarde el proyecto después de aceptar una optimización para conservar el resultado y sus parámetros efectivos.')]),
('08_menu_archivo','Guardar abrir y exportar',
 'El menú Archivo agrupa la persistencia del proyecto y las exportaciones del eje. Guarde una copia del proyecto antes de cambiar de terreno.',
 [('Abrir proyecto','Recupera configuración, extremos y carretera desde un JSON compatible.'),('Guardar proyecto','Guarda el estado actual en el archivo del proyecto.'),('Guardar como','Crea otra copia con el nombre y destino que indique.'),('Exportar CSV y DXF','Genera archivos de intercambio. No sustituyen al proyecto JSON para continuar editando.')]),
('23_guardar_proyecto','Guardar un proyecto que pueda reabrirse',
 'Elija Archivo y Guardar proyecto o Guardar como, indique una carpeta propia y escriba un nombre terminado en .json. Confirme Guardar.',
 [('Carpeta propia','Utilice Documentos o una carpeta de trabajo. Evite guardar sus proyectos dentro de la instalación.'),('Topografía','El proyecto referencia el archivo topográfico; conserve la malla junto con sus respaldos.'),('Reabrir','Use Abrir proyecto y verifique terreno, A/B y parámetros.'),('Archivo anterior','Se mantiene compatibilidad de lectura con proyectos RouteMine; al guardar se usa la identidad RouteTec.')]),
('24_exportar_csv','Exportar la ruta a CSV',
 'Con una ruta calculada, elija Archivo y Exportar ruta CSV. Seleccione carpeta y nombre y confirme Guardar.',
 [('Contenido','El eje se escribe como secuencia ordenada de puntos con coordenadas y medidas del recorrido.'),('Corte y relleno','Las exportaciones del modo de movimiento de tierras incluyen la clasificación de los segmentos.'),('Hoja de cálculo','Abra el CSV con separador coma y punto decimal; conserve la precisión de coordenadas.'),('Comprobación','Verifique los primeros y últimos puntos frente a A y B antes de reutilizarlo en otro programa.')]),
('25_exportar_dxf','Exportar el eje para CAD',
 'Elija Archivo y Exportar eje DXF. Abra el archivo en su programa CAD y compruebe las unidades y coordenadas.',
 [('Eje','El archivo representa el recorrido calculado; la topografía original no se modifica.'),('Capas','Los segmentos de corte y relleno se diferencian cuando corresponde al modo calculado.'),('Coordenadas','Utiliza el mismo sistema de coordenadas y unidades de la entrada.'),('Uso posterior','El eje requiere desarrollo de plataforma, taludes y revisión de ingeniería antes de emplearlo en obra.')]),
('20_menu_ayuda','Consultar detalles sin recargar la pantalla',
 'El menú Ayuda abre los detalles del resultado. Utilícelo para revisar parámetros efectivos, medidas y los motivos registrados durante la búsqueda.',
 [('Detalles del resultado','Muestra información ampliada de la opción seleccionada.'),('Acerca de RouteTec','Consulta la identificación y la versión del aplicativo.'),('Sin ruta','Permite distinguir una búsqueda sin candidato de un presupuesto agotado o una cancelación.'),('Revisar diseño y corte y relleno','Abre la configuración correspondiente desde los detalles. Corte y relleno está disponible en el segundo modo.')]),
('21_detalles_resultado','Interpretar los límites del resultado',
 'Una ruta encontrada supera las comprobaciones implementadas con sus parámetros efectivos. Una búsqueda sin ruta no prueba imposibilidad geométrica.',
 [('Terreno empinado','Puede necesitar rodeos largos. Al 10 %, 100 m de desnivel exigen al menos 1000 m horizontales.'),('Sin solución encontrada','Revise cobertura de la malla, extremos y modo elegido. Después revise precisión y recursos.'),('Corte y relleno','Si su proyecto lo permite, seleccione el segundo modo y establezca límites explícitos. No se aumentan de forma oculta.'),('Alcance','La comprobación del eje no certifica plataforma, peralte, estabilidad de taludes ni maniobrabilidad de un vehículo concreto.')]),
]

def build_user(captures):
    captures=captures or ROOT/'out/release-qa/guide'
    d=setup('Guía de usuario de RouteTec','Uso del aplicativo con procedimientos y capturas de pantalla')
    body(d,'Aprenda a cargar una topografía, definir los extremos, calcular y revisar una carretera, optimizar un tramo y guardar o exportar el resultado. Las imágenes muestran la edición de escritorio. La instalación se explica en la guía independiente.')
    spiral=ROOT/'out/release-qa/ui-spiral/route-iso.png'
    figure(d,spiral if spiral.exists() else captures/'04_ruta_iso.png','RouteTec con un recorrido calculado sobre una superficie sintética de práctica.',4.8)
    chapter(d,'Recorrido de aprendizaje')
    table(d,['Etapa','Qué hará'],[
        ('Preparar','Cargar la topografía y registrar A y B.'),('Diseñar','Revisar pendiente 10 %, radio 30 m y recta 35 m.'),('Calcular','Elegir el modo y supervisar el progreso.'),('Revisar','Inspeccionar colores, cámaras y medidas.'),('Mejorar','Habilitar y optimizar el tramo P1 a P2.'),('Conservar','Guardar JSON y exportar CSV o DXF.')])
    body(d,'Las capturas son sesiones reales del programa y pueden mostrar datos de ejemplo. Los parámetros de cada proyecto se consultan en Configuración; no deduzca los valores iniciales a partir de una ruta abierta.')
    d.add_heading('Atajos de archivo',2)
    table(d,['Atajo','Acción'],[('Ctrl O','Abrir proyecto.'),('Ctrl S','Guardar proyecto.'),('Ctrl Mayús S','Guardar el proyecto como otra copia.')])
    for i,(filename,title,text,rows) in enumerate(USER_PAGES,1):
        chapter(d,f'{i} {title}',text)
        figure(d,captures/(filename+'.png'),f'Figura {i}  {title}',4.5 if len(rows)>3 else 5)
        table(d,['Control o elemento','Uso'],rows)
    for filename,title,explanation in [
      ('ui-spiral/route-iso.png','Ejemplo de ascenso o descenso gradual','En una superficie sintética con laderas del 40 %, RouteTec construye una espiral que distribuye el desnivel sobre un recorrido más largo. En esta sesión el límite de pendiente vial es 10 %, el radio mínimo 30 m y la recta al invertir el giro 35 m.'),
      ('ui-spiral/route-top.png','Seguir las vueltas en planta','La vista superior permite revisar el desarrollo de la espiral y la separación entre vueltas. El terreno de este ejemplo es una cuenca sintética; no representa un levantamiento de campo.'),
      ('ui-features/cut-fill-colors.png','Reconocer corte y relleno en un mismo eje','Ejemplo sintético de clasificación: rojo bajo el terreno, verde sobre el terreno y azul cuando la diferencia está dentro de la tolerancia. La clasificación describe el eje, no el volumen completo de plataforma y taludes.'),
      ('ui-features/optimized-section.png','Examinar el tramo optimizado','La imagen ilustra la selección de un intervalo sobre una ruta sintética y el resultado de su optimización local. Las restricciones del resto del recorrido se conservan.')]:
        extra=ROOT/'out/release-qa'/filename
        if extra.exists():
            chapter(d,title,explanation);figure(d,extra,title,6.0)
    extra=ROOT/'out/release-qa/ui-convex/route-iso.png'
    if extra.exists():
        chapter(d,'Rodear un cerro convexo','Este ejemplo de práctica asciende alrededor de un cerro convexo sintético con laderas del 40 %. Conserva pendiente máxima de 10 %, radio mínimo de 30 m y recta de 35 m al invertir el giro. No necesita corte ni relleno.')
        figure(d,extra,'Recorrido calculado sobre un cerro convexo sintético.',6.1)
    for filename,title in [('route-iso.png','Revisar una carretera en topografía real'),('route-top.png','Seguir el recorrido real en planta')]:
        extra=ROOT/'out/release-qa/real-deep30'/filename
        if extra.exists():
            chapter(d,title,'Ejemplo con Surf_2027_FB.dxf y los extremos del caso revisado. Se calculó en el segundo modo con pendiente máxima 10 %, radio mínimo 30 m, recta 35 m y límites explícitos de corte y relleno de 30 m. Estos 30 m son parámetros de este ensayo y no los valores iniciales de 2 m ni una recomendación de profundidad admisible.')
            figure(d,extra,'Ejemplo real de 10 917,78 m con pendiente máxima comprobada de 9,99 %.',6.0)
    chapter(d,'Tarifas de la configuración avanzada')
    body(d,'Use esta referencia cuando seleccione Costos configurados. El cálculo es preliminar; las tarifas deben proceder de su propio proyecto.')
    table(d,['Campo','Interpretación'],[
      ('Vía por m²','Costo de superficie vial.'),('Corte suelo y corte roca por m³','Tarifas de excavación por tipo de material.'),('Colocación de relleno por m³','Costo de colocar material en relleno.'),('Material importado por m³','Tarifa de suministro externo.'),('Disposición por m³','Costo de retirar o disponer material excedente.'),('Acarreo por m³ km','Tarifa de transporte por volumen y distancia.'),('Fracción de roca','Proporción entre 0 y 1 del material clasificado como roca.'),('Acarreo interno importación y disposición','Distancias utilizadas para estimar el transporte correspondiente.')])
    for filename,caption in [('14b_costos_acarreos.png','Distancias y tarifas de acarreo.'),('14c_perfiles_costos.png','Perfiles disponibles para comparar alternativas.')]:
        extra=captures/filename
        if extra.exists():
            chapter(d,caption.rstrip('.'))
            figure(d,extra,caption,5.7)
    chapter(d,'Resolver problemas frecuentes')
    table(d,['Situación','Acción'],[
      ('No carga la topografía','Compruebe que sea una malla DXF 3DFACE u OBJ y que el archivo exista.'),
      ('Calcular está deshabilitado','Espere a que concluya la carga, registre A/B y revise mensajes de muestreo o recursos.'),
      ('El punto queda fuera del terreno','Vuelva a seleccionarlo sobre una cara visible de la malla.'),
      ('No encuentra ruta','Revise Detalles del resultado, cobertura y restricciones. Aumentar recursos no asegura una solución.'),
      ('La ruta es larga','La pendiente y la geometría obligan a rodeos. Compare el segundo modo solo con límites admisibles.'),
      ('No aparece P1 y P2','Habilite Puntos intermedios desde Configuración. Para marcar necesita una ruta calculada.'),
      ('No puede exportar','Calcule o abra una ruta válida y elija una carpeta donde pueda escribir.'),
      ('Proyecto no recupera topografía','Restaure la malla referenciada desde el respaldo; el JSON no sustituye al archivo del terreno.')])
    body(d,'Organización visual de referencia: manual de usuario de SeqLabel y generador del manual detallado del repositorio DeyviRodrigo/seqlabel. El contenido y las capturas de esta guía corresponden a RouteTec.')
    path=OUT/'Guía de Usuario.docx';d.save(path);return path

def build_installation(captures):
    captures=captures or ROOT/'out/installer-captures'
    d=setup('Guía de instalación de RouteTec','Requisitos y puesta en marcha de la edición Windows de 64 bits')
    body(d,'Instale RouteTec con el paquete RouteTec-Setup-1.0.0.exe. El instalador incluye el aplicativo, Qt, VTK y las bibliotecas de ejecución de Visual C++ necesarias. No necesita instalar un lenguaje de programación, un compilador ni un framework por separado.')
    body(d,'La instalación es local para la cuenta de Windows que la ejecuta. No necesita Internet durante la instalación ni una cuenta de servicio para utilizar el programa.')
    table(d,['Requisito','Condición'],[
      ('Sistema operativo','Windows 10 de 64 bits versión 1809 o posterior, o Windows 11 de 64 bits.'),
      ('Gráficos','Controlador gráfico con OpenGL 3.2 o posterior para la vista tridimensional.'),
      ('Memoria','Se recomiendan 16 GB de RAM para topografías grandes. El consumo depende del terreno y los recursos de búsqueda.'),
      ('Espacio','Reserve al menos 500 MB para aplicativo y paquete. Los levantamientos, proyectos y exportaciones necesitan espacio adicional.'),
      ('Permisos','Permiso de escritura en la carpeta de su perfil. La instalación predeterminada no solicita privilegios de administrador.'),
      ('Programas adicionales','No requiere Python, Visual Studio, conda, Qt ni VTK instalados en el sistema.')])
    body(d,'El ejecutable distribuido utiliza el CPU. No requiere CUDA ni una GPU NVIDIA. La aceleración gráfica de la escena sí depende del controlador OpenGL.')
    pages=[
      ('01-bienvenida','Iniciar el instalador','Cierre una instancia de RouteTec que esté usando la misma carpeta de destino. Haga doble clic en RouteTec-Setup-1.0.0.exe y pulse Siguiente.',[
          ('Archivo','Use el instalador entregado con esta edición. Mantenga una copia para reinstalar.'),('Identidad','El nombre del producto y la versión deben indicar RouteTec 1.0.0.')]),
      ('02-carpeta','Elegir la carpeta de instalación','Conserve la carpeta propuesta o pulse Examinar para seleccionar otro destino propio. Confirme Instalar.',[
          ('Destino inicial',r'%LOCALAPPDATA%\Programs\RouteTec'),('Carpeta de trabajo','Guarde topografías y proyectos en Documentos u otra carpeta independiente de la instalación.')]),
      ('03-archivos-instalados','Comprobar los archivos instalados','El instalador copia el programa, bibliotecas, plugins gráficos, ejemplos y licencias. No descarga componentes de Internet. La captura muestra una instalación de prueba ya completada.',[
          ('Dependencias','Qt y VTK se copian junto al ejecutable.'),('Visual C++','Las bibliotecas de ejecución se incluyen localmente; no se necesita instalar el compilador C++.'),('Interrupción','Si la copia no termina, vuelva a ejecutar el instalador después de cerrar RouteTec.')]),
      ('04_finalizar','Finalizar y abrir RouteTec','Al terminar, pulse Finalizar. Abra RouteTec desde el acceso directo del escritorio o desde el menú Inicio.',[
          ('Inicio','Busque RouteTec en el menú Inicio.'),('Ejecutable','También puede abrir RouteTec.exe en la carpeta instalada.'),('Ejemplos','La subcarpeta Ejemplos contiene mallas de práctica expresadas en metros.')]),
    ]
    for filename,title,text,rows in pages:
        chapter(d,title,text)
        p=captures/(filename+'.png')
        if p.exists():figure(d,p,title,5.2)
        table(d,['Elemento','Qué hacer o comprobar'],rows)
    chapter(d,'Comprobar el primer inicio')
    body(d,'Cargue una malla de la carpeta Ejemplos, registre A y B sobre ella y pulse Calcular ruta. Use un par cercano sobre una zona llana para comprobar primero la apertura, selección y dibujo del resultado. El manual de usuario contiene el procedimiento completo.')
    p=captures/'05_aplicacion_instalada.png'
    if p.exists():figure(d,p,'RouteTec iniciado desde una instalación independiente del entorno de desarrollo.',4.8)
    table(d,['Comprobación','Resultado esperado'],[
      ('Inicio','Aparece RouteTec y la vista tridimensional.'),('Topografía','Se dibuja la malla sin errores de bibliotecas.'),('A y B','Se registran los dos puntos elegidos.'),('Carretera','Se dibuja la ruta cuando el par y las restricciones permiten una solución.'),('Guardar','Puede guardar un JSON en su carpeta de trabajo.')])
    chapter(d,'Usar la edición portable')
    body(d,'Como alternativa al instalador, extraiga RouteTec-Portable-1.0.0.zip en una carpeta propia y abra RouteTec.exe dentro de RouteTec. Mantenga toda la carpeta unida: las DLL, platforms, plugins y Licencias son parte del aplicativo. No ejecute el programa directamente desde la vista comprimida del ZIP.')
    table(d,['Edición','Diferencia práctica'],[
      ('Instalable','Copia los archivos, crea accesos directos y añade un desinstalador.'),('Portable','No necesita instalación. El usuario extrae y conserva completa la carpeta.'),('Funcionalidad','Ambas utilizan el mismo aplicativo y los mismos algoritmos.')])
    chapter(d,'Actualizar y desinstalar')
    body(d,'Antes de actualizar, guarde sus proyectos y cierre RouteTec. Ejecute el instalador de la nueva edición sobre la carpeta de la instalación anterior. Conserve por separado sus topografías y resultados.')
    body(d,'Para desinstalar, abra Configuración de Windows, Aplicaciones, RouteTec y Desinstalar, o ejecute Desinstalar.exe en la carpeta instalada. Confirme únicamente después de guardar y cerrar su trabajo. El desinstalador elimina los archivos de distribución registrados y conserva archivos ajenos al manifiesto. No use la carpeta instalada para almacenar sus únicos respaldos.')
    p=captures/'04-desinstalar.png'
    if p.exists():figure(d,p,'Ventana de confirmación de desinstalación de RouteTec.',4.4)
    chapter(d,'Resolver problemas de instalación e inicio')
    table(d,['Mensaje o síntoma','Acción'],[
      ('Sistema no compatible','Compruebe que Windows sea de 64 bits y cumpla la versión mínima.'),
      ('No puede escribir en destino','Elija una carpeta de su perfil; cierre otras instancias del aplicativo.'),
      ('Falta una DLL','Reinstale el paquete completo. No descargue DLL aisladas de sitios desconocidos.'),
      ('Error del plugin de plataforma Qt','Conserve la carpeta platforms junto al aplicativo; vuelva a extraer o reinstalar el paquete completo.'),
      ('Vista negra o error de OpenGL','Actualice el controlador desde el fabricante de la GPU y compruebe el soporte de OpenGL 3.2. En escritorios remotos depende del controlador disponible.'),
      ('No encuentra un terreno al abrir JSON','Restaure el archivo topográfico referenciado por el proyecto.'),
      ('No encuentra una carretera','La instalación puede ser correcta. Revise A/B, cobertura y restricciones con la guía de usuario.')])
    chapter(d,'Componentes incluidos y requisitos técnicos')
    table(d,['Componente','Función'],[
      ('RouteTec.exe','Aplicativo compilado para Windows x64.'),('Qt 6','Ventanas, controles e integración nativa.'),('VTK','Representación y selección en la escena tridimensional.'),('Visual C++ Runtime','Ejecución del programa C++ y sus bibliotecas. Se distribuye junto al ejecutable.'),('Plugins y dependencias','Soporte de plataforma Windows y bibliotecas utilizadas por Qt y VTK.'),('Licencias','Textos y versiones de terceros incluidos en la carpeta Licencias.')])
    body(d,'No instale lenguajes o frameworks adicionales para ejecutar este paquete. Solo una reconstrucción desde código fuente requiere herramientas de desarrollo; esa tarea queda fuera de la instalación del usuario final.')
    body(d,'Fuentes de requisitos y distribución: documentación oficial de plataformas compatibles de Qt 6.10 y documentación de Microsoft sobre distribución de bibliotecas Visual C++.')
    body(d,'https://doc.qt.io/qt-6.10/supported-platforms.html')
    body(d,'https://learn.microsoft.com/en-us/cpp/windows/redistributing-visual-cpp-files?view=msvc-170')
    path=OUT/'Guía de Instalación.docx';d.save(path);return path

def main():
    p=argparse.ArgumentParser();p.add_argument('kind',choices=['source','user','installation','all']);p.add_argument('--captures',type=Path);a=p.parse_args()
    OUT.mkdir(parents=True,exist_ok=True)
    if a.kind in {'source','all'}: print(build_source())
    if a.kind in {'user','all'}: print(build_user(a.captures))
    if a.kind in {'installation','all'}: print(build_installation(a.captures))

if __name__=='__main__':main()
