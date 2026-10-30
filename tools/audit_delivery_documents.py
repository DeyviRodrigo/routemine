"""Verify source transcription and inventory the rendered delivery documents."""
from pathlib import Path
import hashlib
import json
from docx import Document
import pdfplumber

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'entrega/RouteTec'
QA=ROOT/'out/delivery-docs'

def main():
    manifest=json.loads((QA/'source-manifest.json').read_text(encoding='utf-8'))
    d=Document(OUT/'Código fuente.docx')
    positions={p.text:i for i,p in enumerate(d.paragraphs) if p.style.name=='Heading 1'}
    for f in manifest:
        data=(ROOT/f['path']).read_bytes()
        assert hashlib.sha256(data).hexdigest()==f['sha256'], 'Source changed after generation: '+f['path']
        expected=data.decode('utf-8-sig').splitlines()
        start=positions[f['path']]+3
        actual=[p.text for p in d.paragraphs[start:start+f['lines']]]
        assert actual==expected,'Transcription differs: '+f['path']
    records={'source_files':len(manifest),'source_lines':sum(f['lines'] for f in manifest),'source_transcription':'PASS','documents':[]}
    for p in OUT.glob('*.docx'):
        doc=Document(p);pdf=QA/'renders'/p.stem/(p.stem+'.pdf')
        r={'file':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'pictures':len(doc.inline_shapes),'bytes':p.stat().st_size}
        if pdf.exists():
            with pdfplumber.open(pdf) as rendered:
                r['pages']=len(rendered.pages)
                r['page_text_lengths']=[len((page.extract_text() or '').strip()) for page in rendered.pages]
                assert all(n>40 for n in r['page_text_lengths']),'Blank page: '+p.name
                for page in rendered.pages:
                    # Word wraps long code lines; no body text should escape the physical page.
                    for char in page.chars:
                        assert char['x0']>=-1 and char['top']>=-1 and char['x1']<=page.width+1 and char['bottom']<=page.height+1, ('Clipped character',p.name,page.page_number,char)
                    page.close()
        records['documents'].append(r)
    (QA/'document-audit.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in records.items() if k!='documents'},ensure_ascii=False))
    print('\n'.join(str({k:v for k,v in doc.items() if k!='page_text_lengths'}) for doc in records['documents']))

if __name__=='__main__':main()
